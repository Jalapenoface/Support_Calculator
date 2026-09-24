from io import BytesIO
from datetime import date
from flask import Flask, render_template, request, jsonify, send_file
from scoring import Inputs, calculate
from reportlab.lib import colors
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.enums import TA_CENTER
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, PageBreak
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment

app = Flask(__name__)

TECH = {
  1:"Platform remained stable throughout the semester. Support hours were focused entirely on instructional growth.",
  2:"Minor LTI or link friction detected. Resolved with standard maintenance; minimal impact on student experience.",
  3:"System lag during peak submission windows. Required reactive support and manual interventions to maintain stability.",
  4:"Major tool failure or 1hr+ outage. Significant spike in support tickets and student frustration noted.",
  5:"Critical infrastructure failure (e.g. exam crash). Course performance was severely compromised by platform instability."
}
TEAM = {
  1:"Instructional team is tech-fluent and self-sufficient. Operational overhead for this course is near zero.",
  2:"Team follows standard workflows. Requires routine pre-semester checks and minimal in-flight adjustments.",
  3:"Faculty requires high-touch guidance on Moodle tools. Significant support time allocated to 1-on-1 coaching.",
  4:"High operational demand due to frequent custom requests or last-minute design changes. Taxing on staff bandwidth.",
  5:"Fundamental mismatch in pedagogical approach. Requires intensive management oversight and conflict resolution."
}

def payload():
    d = request.get_json(silent=True) or request.form
    def num(k, default):
        try: return float(d.get(k, default))
        except (TypeError, ValueError): return default
    def integer(k, default):
        try: return int(float(d.get(k, default)))
        except (TypeError, ValueError): return default
    i = Inputs(num('enrollment',120), num('hours',45), num('complexity',6), num('satisfaction',85), num('hourly_rate',35), integer('tech_severity',2), integer('team_severity',1))
    meta = {k: str(d.get(k,'')) for k in ['course_name','faculty_name','department','tech_notes','team_notes']}
    return i, meta

@app.get('/')
def index():
    return render_template('index.html', tech=TECH, team=TEAM)

@app.post('/api/calculate')
def api_calculate():
    i, _ = payload()
    return jsonify(calculate(i))

@app.post('/download/pdf')
def download_pdf():
    i, meta = payload(); r = calculate(i)
    buf = BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=letter, rightMargin=42, leftMargin=42, topMargin=42, bottomMargin=42)
    styles = getSampleStyleSheet()
    title = ParagraphStyle('Title2', parent=styles['Title'], textColor=colors.HexColor('#912338'), fontSize=20, leading=24, alignment=TA_CENTER, spaceAfter=8)
    h2 = ParagraphStyle('H2x', parent=styles['Heading2'], textColor=colors.HexColor('#912338'), spaceBefore=12, spaceAfter=6)
    small = ParagraphStyle('small', parent=styles['BodyText'], fontSize=8.5, leading=11, textColor=colors.HexColor('#475569'))
    story=[Paragraph('eConcordia Course Scorecard', title), Paragraph(f"{meta['course_name'] or 'Course'} | {meta['faculty_name'] or 'Faculty not entered'} | {meta['department'] or 'Department not entered'}", styles['BodyText']), Spacer(1,12)]
    metrics = [
      ['Metric','Result','Interpretation'],
      ['Vitality Grade', f"{r['vitality_grade']} ({r['vitality_score']:.1f})", r['vitality_label']],
      ['Total Support Cost', f"${r['total_cost']:,.2f}", 'Per semester'],
      ['Effort / Student', f"{r['effort_minutes_per_student']:.1f} min", r['effort_label']],
      ['Cost / Student', f"${r['cost_per_student']:,.2f}", r['cost_per_student_label']],
      ['Cost / Satisfied Student', f"${r['cost_per_satisfied_student']:,.2f}", r['cost_per_satisfied_student_label']],
      ['Pedagogical ROI', f"{r['pedagogical_roi']:.1f}%", r['pedagogical_roi_label']],
      ['Scalability Index', f"{r['scalability_index']:.1f}/10", r['scalability_label']],
      ['Efficiency Index', f"{r['efficiency_index']}%", 'Model-derived'],
      ['Resource Utilization', f"{r['resource_utilization']}%", 'Composite indicator'],
      ['Satisfaction', f"{r['satisfaction_score']:.1f}%", 'Entered value'],
    ]
    t=Table(metrics, colWidths=[155,110,220], repeatRows=1)
    t.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,0),colors.HexColor('#912338')),('TEXTCOLOR',(0,0),(-1,0),colors.white),('FONTNAME',(0,0),(-1,0),'Helvetica-Bold'),('FONTSIZE',(0,0),(-1,-1),8.5),('GRID',(0,0),(-1,-1),0.35,colors.HexColor('#cbd5e1')),('VALIGN',(0,0),(-1,-1),'TOP'),('ROWBACKGROUNDS',(0,1),(-1,-1),[colors.white,colors.HexColor('#f8fafc')]),('LEFTPADDING',(0,0),(-1,-1),6),('RIGHTPADDING',(0,0),(-1,-1),6),('TOPPADDING',(0,0),(-1,-1),6),('BOTTOMPADDING',(0,0),(-1,-1),6)]))
    story += [t, Paragraph('Inputs',h2), Paragraph(f"Enrollment: {i.enrollment:g} | Support hours: {i.hours:g} | Complexity: {i.complexity:g}/10 | Satisfaction: {i.satisfaction:g}% | Hourly rate: ${i.hourly_rate:g}", styles['BodyText'])]
    story += [Paragraph('Strategic Recommendations',h2)]
    for a,b in r['recommendations']: story += [Paragraph(f"<b>{a}</b>: {b}", styles['BodyText']), Spacer(1,4)]
    story += [Paragraph('Environmental Context & Risk',h2), Paragraph(f"Context Severity: {r['context_severity']}/10. This is contextual only and does not affect the Vitality Grade.", styles['BodyText']), Spacer(1,6), Paragraph(f"<b>Technical stability:</b> {TECH.get(i.tech_severity,'')}",small), Paragraph(meta['tech_notes'] or 'No technical notes entered.',small), Spacer(1,6), Paragraph(f"<b>Instructional team complexity:</b> {TEAM.get(i.team_severity,'')}",small), Paragraph(meta['team_notes'] or 'No team notes entered.',small)]
    story += [Spacer(1,12), Paragraph('Model note: The scorecard equations and thresholds are operational heuristics reproduced from the supplied prototype. They should be calibrated against eConcordia historical data before being treated as institutional benchmarks.', small)]
    doc.build(story); buf.seek(0)
    name=(meta['course_name'] or 'Course-Audit').replace('/','-').replace('\\','-')
    return send_file(buf, mimetype='application/pdf', as_attachment=True, download_name=f'{name}-{date.today().isoformat()}.pdf')

@app.post('/download/excel')
def download_excel():
    i, meta = payload(); r=calculate(i)
    wb=Workbook(); ws=wb.active; ws.title='Scorecard'
    maroon='912338'; light='F3F4F6'
    ws['A1']='eConcordia Course Scorecard'; ws['A1'].font=Font(size=18,bold=True,color=maroon)
    ws.merge_cells('A1:C1')
    rows=[('Course',meta['course_name']),('Faculty',meta['faculty_name']),('Department',meta['department']),('', ''),('INPUTS',''),('Enrollment',i.enrollment),('Support hours',i.hours),('Complexity',i.complexity),('Satisfaction',i.satisfaction/100),('Hourly rate',i.hourly_rate),('', ''),('RESULTS',''),('Vitality grade',r['vitality_grade']),('Vitality score',r['vitality_score']),('Total support cost',r['total_cost']),('Cost / student',r['cost_per_student']),('Effort minutes / student',r['effort_minutes_per_student']),('Cost / satisfied student',r['cost_per_satisfied_student']),('Pedagogical ROI',r['pedagogical_roi']/100),('Scalability index',r['scalability_index']),('Efficiency index',r['efficiency_index']/100),('Resource utilization',r['resource_utilization']/100),('Context severity',r['context_severity'])]
    for rr,(a,b) in enumerate(rows,start=3): ws.cell(rr,1,a); ws.cell(rr,2,b)
    for rr in [7,14]:
        pass
    for cell in ['A7','A14']:
        ws[cell].font=Font(bold=True,color='FFFFFF'); ws[cell].fill=PatternFill('solid',fgColor=maroon)
    for row in ws.iter_rows(min_row=1,max_row=ws.max_row,min_col=1,max_col=3):
        for c in row: c.alignment=Alignment(vertical='top')
    for c in ['B11','B21','B23','B24']: ws[c].number_format='0.0%'
    for c in ['B12','B17','B18','B20']: ws[c].number_format='$#,##0.00'
    ws.column_dimensions['A'].width=30; ws.column_dimensions['B'].width=22; ws.column_dimensions['C'].width=45
    eq=wb.create_sheet('Equations')
    eq.append(['Metric','Equation','Meaning'])
    equations=[
      ('Total support cost','Hours x Hourly rate','Estimated semester support labour cost.'),
      ('Cost / student','Total cost / Enrollment','Support cost distributed across enrolled students.'),
      ('Effort / student','Hours / Enrollment x 60','Support minutes per enrolled student.'),
      ('Cost / satisfied student','Total cost / (Enrollment x Satisfaction ratio)','Cost adjusted by the estimated satisfied-student count.'),
      ('Pedagogical ROI','Satisfaction / (Complexity x Hours) x 100','Heuristic satisfaction-to-effort ratio; not financial ROI.'),
      ('Scalability index','min(10, Enrollment / (Complexity x Hours) x 10)','Heuristic capacity indicator capped at 10.'),
      ('Efficiency index','min(100, round((Enrollment/Hours) x ((11-Complexity)/10) x 15))','Students-per-support-hour adjusted downward for complexity.'),
      ('Resource utilization','min(100, round(Efficiency x .4 + Satisfaction x .4 + (100-Complexityx10) x .2))','40/40/20 composite of efficiency, satisfaction and inverse complexity.'),
      ('Vitality score','Efficiency x .3 + Satisfaction x .5 + (10-Complexity) x 2','Composite used to assign the letter grade.'),
    ]
    for x in equations: eq.append(x)
    for c in eq[1]: c.font=Font(bold=True,color='FFFFFF'); c.fill=PatternFill('solid',fgColor=maroon)
    eq.column_dimensions['A'].width=26; eq.column_dimensions['B'].width=65; eq.column_dimensions['C'].width=65
    for row in eq.iter_rows():
        for c in row: c.alignment=Alignment(vertical='top',wrap_text=True)
    out=BytesIO(); wb.save(out); out.seek(0)
    name=(meta['course_name'] or 'Course-Audit').replace('/','-').replace('\\','-')
    return send_file(out, mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet', as_attachment=True, download_name=f'{name}-{date.today().isoformat()}.xlsx')

if __name__ == '__main__':
    app.run(debug=True)
